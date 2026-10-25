"""Independent source/report reconciliation and interactive dashboard checks."""
import hashlib, math
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from analysis import ROOT, SOURCE, load, tables

source=load_workbook(SOURCE,data_only=True)
rows=list(source['Form Responses 1'].values); values=rows[1:]
assert len(values)==560
report=load_workbook(ROOT/'reports/feedback_analysis.xlsx',data_only=True)
k=dict(list(report['KPIs'].values)[1:])
for col,key in [(12,'course_rating'),(18,'teaching_rating'),(25,'mentor_rating'),(35,'expectations_rating')]:
    expected=sum(r[col] for r in values)/len(values)
    assert math.isclose(k[key],expected,rel_tol=1e-12),(key,k[key],expected)
assert math.isclose(k['continuation_pct'],100*sum(r[31]=='جی ہاں، ضرور' for r in values)/len(values))
assert Counter(r[2] for r in values)==dict((r[0],r[1]) for r in list(report['Campus Analysis'].values)[1:])
assert list(report['Raw Responses'].values)==rows
assert len(report['Charts']._charts)==3
assert report['Correlations'].conditional_formatting
original=Path('/workspace/attachments/63b2c67f-a62e-4eb1-9b8e-6a66a4600aae/Ilm-e-Deen Course 2026 - Mid-Session Student Feedback Form (Responses).xlsx')
assert hashlib.sha256(original.read_bytes()).digest()==hashlib.sha256(SOURCE.read_bytes()).digest()
raw,d=load(); assert len(d)==560 and tables(raw,d)['Rating Statistics'].invalid.sum()==0
app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=60).run()
assert not app.exception, app.exception
assert app.metric[0].value=='560'
app.sidebar.multiselect[0].select('Lahore').run()
assert not app.exception
assert app.metric[0].value=='81'
app.sidebar.multiselect[1].set_value(['NONEXISTENT']).run()
assert not app.exception
assert len(app.warning)>0
print('PASS: independent raw-cell KPI calculations; campus counts; raw sheet preservation; source SHA256; 3 Excel charts; heatmap; dashboard default, campus filter, and empty-state behavior.')
